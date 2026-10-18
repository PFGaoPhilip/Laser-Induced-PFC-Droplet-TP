"""English scientific redraws for Chapter 1 and Chapter 2, using SCI-PPT native scenes."""
import sys,json,math,re,textwrap
from pathlib import Path
from html import escape
ROOT=Path(__file__).resolve().parents[1]
SKILL=Path.home()/'.codex/skills/SCI-PPT/scripts'
sys.path.insert(0,str(SKILL))
import render_scene
from pptx import Presentation
from pptx.util import Inches
OUT=ROOT/'assets/figures-editable/c1-c2'
BG='#142632'; FG='#E7F0F3'; ACC='#8EDBD4'; MUT='#B4C7D2'; AREA='#213E4B'; HOT='#F6BC77'
class Canvas:
 def __init__(self,eq):
  self.eq=eq;self.o=[];self.n=0
  source=ROOT/'assets/figures-editable/source-svg'/f"{eq['id'].lower()}.svg"
  if not source.is_file():source=ROOT/'assets/figures'/f"{eq['id'].lower()}.svg"
  old=source.read_text(encoding='utf8')
  self.h=int(re.search(r'viewBox="0 0 980 (\d+)',old)[1]);self.h=max(self.h,650)
  self.text(26,22,eq['id']+' | Physical-variable diagram',19,ACC,925)
 def add(self,t,**kw):
  self.n+=1;self.o.append(dict(type=t,id=f"{self.eq['id']}-{self.n:03}",name=kw.pop('name',t+str(self.n)),**kw))
 def text(self,x,y,s,size=19,color=FG,w=400):
  for a,b in [('d2R/dt2','R̈'),('dR/dt','Ṙ'),('p_inf','p∞'),('sigma','σ'),('rho','ρ'),('tau','τ'),('delta_T','δ_T'),('Delta','Δ'),('Pi','Π'),('4pi','4π'),('2sigma','2σ'),('Qdot','Q̇'),('Vdot','V̇'),('R2','R²'),('r2','r²'),('>=','≥'),('<=','≤'),('<<','≪'),('->','→')]:s=s.replace(a,b)
  s=re.sub(r'(?<![A-Za-z])mu(?![A-Za-z])','μ',s)
  self.add('text',x=x/75,y=y/75,w=w/75,h=(size*1.55)/75,text=s,font='Arial',font_size=size*.96,color=color,word_wrap=False)
 def rect(self,x,y,w,h,fill=AREA,line=MUT):self.add('rect',x=x/75,y=y/75,w=w/75,h=h/75,fill=fill,line=line,line_width=1.5)
 def circle(self,x,y,r,fill=AREA,line=MUT):self.add('ellipse',x=(x-r)/75,y=(y-r)/75,w=2*r/75,h=2*r/75,fill=fill,line=line,line_width=1.7)
 def line(self,x,y,X,Y,arrow=False,color=ACC):self.add('arrow' if arrow else 'line',x1=x/75,y1=y/75,x2=X/75,y2=Y/75,line=color,line_width=2)
 def curve(self,pts,color=ACC):
  for a,b in zip(pts,pts[1:]):self.line(*a,*b,color=color)
 def box(self,x,y,w,h,txt,fill=AREA):self.rect(x,y,w,h,fill);self.text(x+10,y+10,txt,18,w=w-15)
 def axes(self,x=570,y=325,w=320,h=205,xlabel='Coordinate',ylabel='Response'):
  self.line(x,y,x+w,y,True);self.line(x,y,x,y-h,True);self.text(x+w-80,y+10,xlabel,17,w=150);self.text(x-10,y-h-30,ylabel,17,w=300)
 def sphere(self,liquid=False):
  self.circle(270,215,100);self.line(270,215,370,215,True);self.text(285,175,'a' if liquid else 'R(t)',21,w=130)
  self.text(220,230,'PFC liquid' if liquid else 'p_b',20,w=180)
  self.text(430,125,'Aqueous carrier' if liquid else 'Surrounding liquid',22,w=400)
  self.text(430,163,'p_c; interface: sigma_pc, shell' if liquid else 'p_inf, rho, mu; interface: sigma',19,w=480)
  if not liquid:
   self.line(270,215,410,310,True);self.text(440,300,'r >= R',18,w=160)
   self.line(410,215,475,215,True);self.text(415,235,'+e_r (coordinate)',17,w=270)
   self.text(100,375 if self.eq['id']=='C1-E02' else 345,'Signed wall speed: dR/dt; acceleration: d2R/dt2',18,w=800)
  else:self.text(100,345,'Finite liquid core: a differs from subsequent bubble radius R.',18,w=800)
 def finish(self):
  labels=self.eq['diagram'].get('labels',[]);rows=[[],[]]
  for i,label in enumerate(labels):
   if isinstance(label,list):label=label[0]
   rows[i%2]+=textwrap.wrap(label,51,break_long_words=False,break_on_hyphens=False)
  start=408
  for col,lines in enumerate(rows):
   for j,line in enumerate(lines):self.text(30+col*475,start+j*24,line,17,w=455)
  yy=max(start+max(map(len,rows))*24+16,515)
  for note in self.eq['diagram'].get('notes',[]):
   if isinstance(note,list):note=note[0]
   for ln in textwrap.wrap(note,102,break_long_words=False):self.text(30,yy,ln,16,MUT,925);yy+=24
  self.h=max(self.h,yy+25)
  return {'slide':{'width':980/75,'height':self.h/75},'background':BG,'objects':self.o}

def rayleigh_points():
 # x = R/Rmax = 1-s^2 regularizes the rest-at-maximum endpoint.
 # Integrate dt/ds proportional to 2s*x^(3/2)/sqrt(1-x^3).
 N=2000; ss=[i/N for i in range(N+1)]; f=[]
 for z in ss:
  x=1-z*z
  f.append(2/math.sqrt(3) if z==0 else 2*z*x**1.5/math.sqrt(1-x**3))
 tt=[0.0]
 for i in range(N):tt.append(tt[-1]+(f[i]+f[i+1])/(2*N))
 return [(605+280*tt[i]/tt[-1],145+165*ss[i]**2) for i in range(0,N+1,50)]

def draw(eq):
 c=Canvas(eq);id=eq['id'];k=eq['diagram']['type'];n=int(id.split('E')[1]);ch=id[:2];fix=[]
 # Numerical drawing cases retain their original physical ordering.
 if ch=='C1' and n>=12:n+=1
 if ch=='C1' and k=='sphere':
  c.sphere();fix+=['Outward arrow explicitly labels radial coordinate, not signed wall motion.']
  if n in [13,14,15,19,33]:c.line(160,215,205,215,True,color=HOT);c.text(90,80,'Collapse: dR/dt < 0',18,HOT,350);fix+=['Separate inward collapse arrow from outward radial coordinate.']
  if n==2:
   c.circle(270,215,150,None);c.text(480,265,'u(2R) = (dR/dt)/4',20,w=400)
  if n==4:c.circle(600,300,5,HOT);c.text(615,330,'Fixed r',18,HOT,200);c.text(100,85,'Moving interface versus fixed observer',19,w=650)
  if n==6:c.text(450,275,'p_l = p_b - 2sigma/R - 4mu(dR/dt)/R',18,w=470)
  if n==26:
   c.line(80,95,80,340,color=MUT);c.line(80,100,270,100);c.text(130,70,'h',18,w=120);c.line(135,215,170,215,True);c.text(480,270,'Boundary type must be specified.',18,w=480);fix+=['Stand-off h reaches center from an unspecified boundary.']
 elif ch=='C1' and k=='impulse':
  if n==21:
   c.axes(160,325,650,210,'t','p - p_ref');pts=[(160,325),(255,325),(255,160),(510,160),(510,325),(810,325)];c.curve(pts);c.rect(256,162,252,160,AREA,None);c.text(295,210,'Area = Pi(x)',22,w=320);c.text(240,335,'t0',18,w=80);c.text(495,335,'t1',18,w=80)
  else:
   c.rect(100,180,680,120);c.line(100,180,780,180,color=MUT);c.line(100,300,780,300,color=MUT);c.line(190,240,650,240,True);c.text(115,120,'Pi0 at z = 0',20,w=300);c.text(600,120,'Pi = 0 at z = L',20,w=350);c.text(255,245,'Delta u_z toward outlet',19,w=400);c.text(185,335,'Pressure impulse gradient drives early acceleration.',19,w=740)
  fix+=['Pressure-time integral or spatial impulse gradient distinguished explicitly.']
 elif ch=='C1' and k=='jet':
  c.rect(110,125,750,200,None);c.circle(285,225,65);c.line(350,225,680,225,True);c.text(530 if n==28 else 420,135 if n==28 else 170,'Liquid u(x,t), p(x,t)',21,w=440);c.text(630,280,'Outlet',19,w=160)
  if n==28:
   c.line(350,225,305,225,True,color=HOT);c.text(330,250,'n: liquid -> gas',18,HOT,320);c.line(350,225,350,170,True);c.text(360,175,'s: tangent',18,w=200)
  c.text(130,345,'Moving interface Gamma; spatial field required.',19,w=800);fix+=['Domain, moving interface and normal distinguished; no radial model claims a jet.']
 elif ch=='C1' and k=='timescale':
  if n in [16,17]:
   c.circle(200,205,85);c.circle(440,205,32);c.line(300,205,390,205,True);c.text(125,310,'R_max, initially at rest',18,w=300);c.text(390,280,'R decreases',18,w=200);c.axes(600,315,290,190,'t','R');c.curve(rayleigh_points());c.text(635,85,'Ideal collapse trajectory',17,w=300)
  else:
   c.box(100,150,300,95,'Wall speed / sound speed');c.box(570,150,300,95,'Propagation / event time');c.line(405,198,560,198,True);c.text(120,300,'Check both motion and pressure communication.',20,w=790)
  fix+=['Qualitative trajectories have no fabricated data points.']
 elif ch=='C1' and id=='C1-E08':
  c.circle(270,215,100);c.text(170,190,'PFC vapor + water vapor',18,w=350);c.text(180,230,'+ permanent gas',18,w=310);c.circle(305,275,20, HOT);c.text(450,180,'p_b = p_v,PFC + p_v,w + p_g',21,w=480);c.text(450,240,'Residual PFC liquid: finite inventory',18,w=480);fix+=['Residual liquid kept distinct from vapor mixture.']
 elif ch=='C1' and id=='C1-E31':
  c.box(60,160,235,100,'Finite PFC / heat input');c.box(370,160,235,100,'Bubble pressure p_b');c.box(680,160,235,100,'Carrier displacement');c.line(300,210,360,210,True);c.line(610,210,670,210,True);c.text(130,320,'Geometry selects direction; scalar radius does not.',20,w=780)
 elif ch=='C1' and k=='energy':
  if n==18:
   c.box(50,165,250,105,'Pressure-work budget E_B');c.box(370,165,250,105,'Liquid kinetic energy K_l');c.box(690,165,240,105,'Remaining volume potential');c.line(305,215,360,215,True);c.text(100,320,'K_l / E_B = 1 - (R/R_max)^3; ideal loss-free control.',19,w=850)
  elif n in [9,12]:
   c.circle(280,215,80);c.circle(280,215,140,None);c.circle(280,215,155,None);c.text(200,210,'Cavity',22,w=200);c.line(280,215,425,215,True);c.text(450,160,'Exterior liquid shells: 4pi r2 dr',20,w=500);c.text(450,215,'Integrate r from R to infinity',19,w=480);c.text(450,275,'Storage K_l' if n==9 else 'Positive dissipation density',20,w=460)
  else:
   c.box(50,165,250,105,'Pressure work');c.box(370,165,250,105,'K_l + E_sigma');c.box(690,165,240,105,'Viscous dissipation');c.line(305,215,360,215,True);c.line(625,215,680,215,True);c.text(100,320,'Storage can change sign; dissipation is nonnegative.',20,w=850)
  fix+=['Kinetic energy resides outside cavity; surface storage and dissipation separated.']
 elif ch=='C2' and n in [1,2,3]:
  c.circle(220,215,105,None);c.circle(220,215,63,None);c.line(220,215,325,215,True);c.text(240,225,'w',20,w=80);c.text(115,335,'Illuminated transverse plane',18,w=400);c.axes(540,325,350,210,'r / w','F / F0');c.curve([(545+i*340/40,325-200*math.exp(-2*(i/40*2)**2)) for i in range(41)]);c.text(595,85,'Gaussian profile (prescribed)',18,w=350)
  if n==3:c.circle(220,215,50,None,HOT);c.text(85,70,'Centered aperture b_a',19,HOT,350)
  fix+=['Transverse Gaussian distinguished from propagation depth; beam radius defined at exp(-2).']
 elif ch=='C2' and n==4:
  c.axes(160,330,650,220,'t','f(t)');c.curve([(165,330),(290,330),(290,140),(620,140),(620,330),(800,330)]);c.text(315,195,'Area = 1',24,w=300);c.text(340,345,'Pulse duration tau_L',19,w=350);fix+=['Normalized temporal waveform shown rather than spatial beam.']
 elif ch=='C2' and n in [5,6,32]:
  c.rect(310,110,260,225);c.line(90,205,300,205,True,color=HOT);c.line(575,205,790,205,True);c.line(310,205,190,115,True);c.text(70,240,'Incident light',19,w=230);c.text(70,80,'Reflection',18,w=230);c.text(335,250,'Absorber: Q_abs',19,w=250);c.text(600,240,'Transmitted light',19,w=300);c.line(330,310,550,310,True);c.text(370,330,'+z; thickness h_a',18,w=350);fix+=['Reflection, absorption and transmission separated; z points into absorber.']
 elif ch=='C2' and n in [7,8,9,10]:
  c.rect(100,125,250,200,HOT);c.rect(350,125,400,200);c.text(125,150,'Absorber / side A',21,BG,300);c.text(465,150,'PFC / side B',21,w=300);c.line(250,235,520,235,True);c.text(380,260,'q; normal A -> B',19,w=400);c.text(100,90,'T_A',20,w=150);c.text(580,90,'T_B',20,w=150);c.text(115,345,'Thermal path and contact are distinct from optical deposition.',18,w=850)
  if n in [9,10]:c.rect(351,126,18,198,HOT,None);c.text(770,200,'delta_T',18,w=160);c.text(770,235,'<< L_h',18,w=160)
  fix+=['Heat transport shown across actual thermal path, without implying uniform core heating.']
 elif ch=='C2' and n in [11,27]:c.sphere(True);fix+=['Finite PFC liquid core uses a or a0, p_d and outer-interface sigma_pc.']
 elif ch=='C2' and n==30:
  c.circle(180,215,26,HOT);c.circle(650,215,140);c.line(230,215,485,215,True);c.text(100,270,'Liquid: a0 = 5 um',19,w=300);c.text(490,365,'Vapor: R_b,inv = 26.6827 um',18,w=450);c.text(90,80,'Same PFC mass',19,w=620);fix+=['Liquid-to-vapor radius ratio shown approximately 5.34, with inventory radius distinct from maximum.']
 elif ch=='C2' and n in [12,13]:
  c.box(80,140,300,100,'Liquid | vapor equilibrium');c.axes(520,325,365,200,'Temperature','Saturation pressure');c.curve([(525+i*350/30,315-180*(i/30)**2) for i in range(31)]);c.text(80,285,'Schematic trend within stated fit ranges',18,w=430);c.text(80,250,'Not an activation threshold',19,HOT,400);fix+=['Equilibrium vapor-pressure trend marked schematic; no activation guarantee.']
 elif ch=='C2' and n in [14,15,16,33]:
  c.circle(225,220,115);c.circle(225,220,29,HOT);c.text(140,70,'Finite PFC liquid: p_d',19,w=400);c.text(110,350,'Inner interface sigma_vp; outer sigma_pc',17,w=460);c.axes(550,320,350,200,'r_n','W');c.curve([(555+i*330/40,320-100*(3*(i/40*1.5)**2-2*(i/40*1.5)**3)) for i in range(41)]);c.text(705,180,'r*, W*',20,w=180);fix+=['New vapor nucleus shown inside finite PFC; barrier maximum distinguished from stable equilibrium.']
 elif ch=='C2' and n in [17,34]:
  c.box(90,140,320,110,'Local rate J(T,p) in V_d');c.line(420,195,510,195,True);c.axes(550,325,350,200,'Integrated hazard','P_act');c.curve([(555+i*330/35,320-185*(1-math.exp(-i/35*3))) for i in range(36)]);c.text(110,300,'Integrate the local rate over space and time.',18,w=780);fix+=['Stochastic activation probability shown versus integrated hazard, not a nucleation barrier curve.']
 elif ch=='C2' and n in [18,35]:
  c.box(80,140,290,110,'Initial PFC mass m_PFC,0');c.line(380,195,480,195,True)
  for i,t in enumerate(['Liquid m_l','Vapor m_v','Dissolved m_diss','Escaped m_esc']):c.box(510+(i%2)*220,100+(i//2)*130,205,95,t)
  c.text(90,355,'Sum of all compartments equals initial inventory.',20,w=850);fix+=['Mass compartments explicitly conserved; dissolved and escaped species included.']
 elif ch=='C2' and n==19:
  c.circle(260,215,115);c.circle(320,260,38,HOT);c.line(330,230,410,145,True);c.text(440,135,'Actual PFC contact patch Gamma_s',20,w=470);c.text(450,220,'Full 4pi R2 requires complete contact.',18,w=470);c.text(130,350,'Bubble boundary and PFC evaporation surface can differ.',19,w=800);fix+=['PFC evaporation localized to actual contacted surface, not entire bubble.']
 elif ch=='C2' and n in [20,21,36]:
  c.rect(80,120,350,220);c.rect(430,120,430,220,None);c.text(180,145,'Vapor',23,w=200);c.text(600,145,'Liquid',23,w=200);c.line(610,230,300,230,True,color=HOT);c.text(480,255,'n: liquid -> vapor',20,HOT,340);c.line(430,310,545,310,True);c.text(450,345,'+e_r, wall motion',18,w=350);c.text(100,85,'Interface Gamma',19,w=300);fix+=['Liquid-to-vapor normal points into cavity; coordinate normal and wall motion kept separate.']
 elif ch=='C2' and n in [24,25]:
  for i,t in enumerate(['Species masses m_v,s','Temperature T_b','Bubble volume V_b']):c.box(60+i*310,120,280,95,t);c.line(200+i*310,225,490,305,True)
  c.box(310,295,360,65,'Mixture pressure p_b');fix+=['Mass, temperature and volume independently control mixture pressure.']
 elif ch=='C2' and n==26:
  c.axes(180,325,650,200,'Vapor volume V_b','PFC partial pressure');c.curve([(185,165),(420,165),(450,165),(500,205),(590,250),(700,278),(810,293)]);c.text(210,120,'Saturation branch',20,w=300);c.text(500,175,'Finite-mass branch: p ~ 1/V',18,w=440);fix+=['Saturation plateau transitions to finite-inventory pressure decay.']
 elif ch=='C2' and k=='energy':
  if n in [22,23]:
   c.box(330,160,300,110,'Bubble energy storage U_b');c.line(100,210,320,210,True);c.text(90,125,'Heat Qdot_b',19,w=230);c.line(480,100,480,150,True);c.text(410,65,'Mass enthalpy entry',19,w=350);c.line(640,215,860,215,True);c.text(690,145,'Work p_b Vdot_b',19,w=270);c.text(110,335,'Conductive heat and transported enthalpy are separate terms.',18,w=820)
  elif n==31:
   c.box(70,150,250,100,'Incident optical energy');c.box(365,150,250,100,'Absorbed: 223.848 nJ');c.box(660,150,250,100,'PFC prep: 97.8243 nJ');c.line(325,200,355,200,True);c.line(620,200,650,200,True);c.text(100,315,'Required thermal-transfer fraction: at least 0.437012.',19,w=850)
  elif n==29:
   c.box(75,150,250,100,'Liquid heating');c.box(365,150,250,100,'Latent conversion');c.box(655,150,250,100,'Vapor heating');c.text(330,185,'+',24,w=40);c.text(620,185,'+',24,w=40);c.text(100,315,'Sum along one constant-pressure path gives Q_iso.',19,w=850)
  else:
   c.box(75,150,250,100,'Sensible heating');c.box(365,150,250,100,'Latent conversion');c.box(655,150,250,100,'Declared preparation');c.text(330,185,'+',24,w=40);c.text(620,185,'=',24,w=40);c.text(100,315,'Path estimate: retain separate loss and mechanical-work budgets.',18,w=850)
  fix+=['Energy ledger separates heat, stored energy, mass enthalpy and pressure work.']
 else:raise ValueError('Unhandled '+id)
 return c.finish(),fix

def svg(scene,eq):
 W=980;H=scene['slide']['height']*75
 p=[f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H:g}" role="img"><title>{escape(eq["id"]+" "+eq["caption_en"])}</title><defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0 0L8 4L0 8Z" fill="{ACC}"/></marker></defs>',f'<rect width="980" height="{H:g}" fill="{BG}"/>']
 for o in scene['objects']:
  t=o['type'];st=f'fill="{o.get("fill") or "none"}" stroke="{o.get("line") or "none"}" stroke-width="{o.get("line_width",1)*1.3}"'
  if t=='text':p.append(f'<text x="{o["x"]*75+2:g}" y="{o["y"]*75+o["font_size"]/.96*1.08:g}" font-family="Arial,sans-serif" font-size="{o["font_size"]/.96:g}" fill="{o["color"]}">{escape(o["text"])}</text>')
  elif t=='rect':p.append(f'<rect x="{o["x"]*75:g}" y="{o["y"]*75:g}" width="{o["w"]*75:g}" height="{o["h"]*75:g}" {st}/>')
  elif t=='ellipse':p.append(f'<ellipse cx="{(o["x"]+o["w"]/2)*75:g}" cy="{(o["y"]+o["h"]/2)*75:g}" rx="{o["w"]*37.5:g}" ry="{o["h"]*37.5:g}" {st}/>')
  else:p.append(f'<line x1="{o["x1"]*75:g}" y1="{o["y1"]*75:g}" x2="{o["x2"]*75:g}" y2="{o["y2"]*75:g}" {st}'+(' marker-end="url(#arrow)"' if t=='arrow' else '')+'/>')
 p.append('</svg>');return '\n'.join(p)
def main():
 OUT.mkdir(parents=True,exist_ok=True);d=json.loads((ROOT/'verification/equations.json').read_text(encoding='utf8'));audit=[]
 for eq in d:
  if not eq['id'].startswith(('C1-','C2-')):continue
  scene,fix=draw(eq);stem=eq['id'].lower();(OUT/(stem+'.scene.json')).write_text(json.dumps(scene,ensure_ascii=False,indent=2),encoding='utf8');(OUT/(stem+'.svg')).write_text(svg(scene,eq),encoding='utf8');render_scene.render(scene,OUT/(stem+'.pptx'))
  audit.append({'id':eq['id'],'corrections':fix,'source':'verification/equations.json','science_status':'Checked against adjacent lesson equations; schematic, not simulated data','visual_status':'pending actual PPTX rendering'})
 # A common canvas preserves native geometry and relative scales across slides.
 for chapter in ['c1','c2']:
  files=sorted(OUT.glob(chapter+'-*.scene.json')); scenes=[json.loads(f.read_text(encoding='utf8')) for f in files]
  deck=Presentation();deck.slide_width=Inches(980/75);deck.slide_height=Inches(max(s['slide']['height'] for s in scenes))
  for scene in scenes:
   slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=render_scene.rgb(BG)
   for o in scene['objects']:render_scene.render_object(slide,o)
  deck.save(OUT/(chapter+'-figures.pptx'))
  (OUT/(chapter+'-slide-order.json')).write_text(json.dumps([f.name.split('.')[0].upper() for f in files]),encoding='utf8')
 (OUT/'audit.json').write_text(json.dumps(audit,indent=2),encoding='utf8');print(f'Wrote {len(audit)} editable figures')
if __name__=='__main__':main()
