"""Shared native SCI-PPT scene, SVG and editable-deck utilities.

Numerical SVG import preserves source curve coordinates without solving models.
"""
from pathlib import Path
from html import escape
from copy import deepcopy
import importlib.util
import json
import math
import re
import sys
import xml.etree.ElementTree as ET

SKILL = Path.home() / '.codex/skills/SCI-PPT/scripts'
if str(SKILL) not in sys.path:
    sys.path.insert(0, str(SKILL))
import render_scene as sci
from scene_support import validate_scene
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.xmlchemy import OxmlElement
from fontTools.misc.transform import Transform
from fontTools.pens.recordingPen import RecordingPen
from fontTools.pens.transformPen import TransformPen
from fontTools.svgLib.path import parse_path


def flatten_objects(objects):
    for obj in objects:
        if obj['type'] == 'group':
            yield from flatten_objects(obj['objects'])
        else:
            yield obj


def _finish_styles(slide, scene):
    """Keep native dashed lines when importing quantitative chart series."""
    styles = {obj.get('name', obj.get('id')): obj for obj in flatten_objects(scene['objects'])}
    for shape in slide.shapes:
        obj = styles.get(shape.name)
        if obj and obj.get('_open'):
            for close in shape._element.xpath('.//a:custGeom/a:pathLst/a:path/a:close'):
                close.getparent().remove(close)
        if obj and obj.get('_dash') and hasattr(shape, 'line'):
            dash = OxmlElement('a:prstDash')
            dash.set('val', 'dash')
            shape.line._get_or_add_ln().append(dash)


def render_scene(scene, path):
    """Use the installed SCI-PPT native scene renderer and reopen the result."""
    path = Path(path)
    sci.render(scene, path)
    if any(obj.get('_dash') or obj.get('_open') for obj in flatten_objects(scene['objects'])):
        prs = Presentation(path)
        _finish_styles(prs.slides[0], scene)
        prs.save(path)
    return path


def make_deck(scenes, path):
    """Render every scene through SCI-PPT; pad differing aspects without distortion."""
    scenes = list(scenes)
    if not scenes:
        raise ValueError('Empty scene list')
    if isinstance(scenes[0], (str, Path)):
        inputs=[Presentation(p) for p in scenes]
        prs=Presentation()
        prs.slide_width=max(p.slide_width for p in inputs)
        prs.slide_height=max(p.slide_height for p in inputs)
        for presentation in inputs:
            for original in presentation.slides:
                slide=prs.slides.add_slide(prs.slide_layouts[6])
                original_bg=original._element.cSld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
                if original_bg is not None:
                    slide._element.cSld.insert(0,deepcopy(original_bg))
                for shape in original.shapes:
                    slide.shapes._spTree.insert_element_before(deepcopy(shape.element),'p:extLst')
        path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);prs.save(path)
        return path
    prs = Presentation()
    width = max(float(s.get('slide', {}).get('width', 13.333)) for s in scenes)
    height = max(float(s.get('slide', {}).get('height', 7.5)) for s in scenes)
    prs.slide_width, prs.slide_height = Inches(width), Inches(height)
    for scene in scenes:
        validate_scene(scene)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = sci.rgb(scene.get('background', '#FFFFFF'))
        for obj in scene['objects']:
            sci.render_object(slide, obj)
        _finish_styles(slide, scene)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    reopened = Presentation(path)
    assert len(reopened.slides) == len(scenes)
    return path


def scene_to_svg(scene, path=None):
    """Export matching scene geometry with live SVG text and embedded styling."""
    validate_scene(scene)
    width = float(scene.get('slide', {}).get('width', 13.333)) * 96
    height = float(scene.get('slide', {}).get('height', 7.5)) * 96
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width:g} {height:g}" role="img">',
             '<title>' + escape(scene.get('_title', 'Scientific schematic')) + '</title>',
             '<defs><marker id="head" markerWidth="7" markerHeight="7" refX="6" refY="3.5" orient="auto-start-reverse" markerUnits="strokeWidth"><path d="M0,0 L7,3.5 L0,7 Z" fill="context-stroke"/></marker></defs>',
             f'<rect width="{width:g}" height="{height:g}" fill="{scene.get("background", "#FFFFFF")}"/>']
    def draw(obj, inherited=1):
        kind = obj['type']
        opacity = float(obj.get('opacity', 1)) * inherited
        if kind == 'group':
            for child in obj['objects']:
                draw(child, opacity)
            return
        get = lambda key, default=0: float(obj.get(key, default)) * 96
        ident = escape(str(obj.get('id', obj.get('name', ''))), quote=True)
        fill = obj.get('fill', '#FFFFFF') or 'none'
        stroke = obj.get('line', '#333333') or 'none'
        style = f'fill="{fill}" stroke="{stroke}" stroke-width="{float(obj.get("line_width", 1.25))*96/72:g}" opacity="{opacity:g}"'
        if obj.get('_dash'):
            style += ' stroke-dasharray="6 4"'
        rotation = float(obj.get('rotation', 0))
        transform = f' transform="rotate({rotation:g} {get("x")+get("w")/2:g} {get("y")+get("h")/2:g})"' if rotation else ''
        if kind == 'text':
            size = float(obj.get('font_size', 18)) * 96 / 72
            x, y, w, h = [get(k) for k in ('x', 'y', 'w', 'h')]
            lines = str(obj['text']).split('\n')
            align = obj.get('align', 'left')
            anchor = {'left': 'start', 'center': 'middle', 'right': 'end'}[align]
            margin = float(obj.get('margin', .03))*96
            tx = x + (w/2 if align == 'center' else w-margin if align == 'right' else margin)
            valign = obj.get('valign', 'middle')
            block_h = size * 1.18 * len(lines)
            base = y + margin + size*.82 if valign == 'top' else y+h-block_h+size*.82-margin if valign == 'bottom' else y+(h-block_h)/2+size*.82
            color = obj.get('color', '#222222')
            font = escape(obj.get('font', 'Arial'), quote=True)
            attrs = f'font-family="{font},sans-serif" font-size="{size:g}" fill="{color}" text-anchor="{anchor}"'
            if obj.get('bold'): attrs += ' font-weight="bold"'
            if obj.get('italic'): attrs += ' font-style="italic"'
            parts.append(f'<text id="{ident}" {attrs} opacity="{opacity:g}"{transform}>')
            for index, line in enumerate(lines):
                parts.append(f'<tspan x="{tx:g}" y="{base+index*size*1.18:g}">{escape(line)}</tspan>')
            parts.append('</text>')
        elif kind in {'line', 'arrow', 'dimension_arrow'}:
            marker = ' marker-end="url(#head)"' if kind != 'line' and obj.get('end_arrowhead', obj.get('arrowhead', 'triangle')) != 'none' else ''
            if obj.get('start_arrowhead', 'none') != 'none': marker += ' marker-start="url(#head)"'
            parts.append(f'<line id="{ident}" x1="{get("x1"):g}" y1="{get("y1"):g}" x2="{get("x2"):g}" y2="{get("y2"):g}" {style}{marker}/>')
        elif kind in {'rect', 'round_rect'}:
            radius = f' rx="{min(get("w"), get("h"))*.08:g}"' if kind == 'round_rect' else ''
            parts.append(f'<rect id="{ident}" x="{get("x"):g}" y="{get("y"):g}" width="{get("w"):g}" height="{get("h"):g}"{radius} {style}{transform}/>')
        elif kind == 'ellipse':
            parts.append(f'<ellipse id="{ident}" cx="{get("x")+get("w")/2:g}" cy="{get("y")+get("h")/2:g}" rx="{get("w")/2:g}" ry="{get("h")/2:g}" {style}{transform}/>')
        elif kind == 'polygon':
            points = ' '.join(f'{float(x)*96:g},{float(y)*96:g}' for x,y in obj['points'])
            tag = 'polyline' if obj.get('_open') else 'polygon'
            parts.append(f'<{tag} id="{ident}" points="{points}" {style}/>')
        elif kind == 'triangle':
            x,y,w,h = [get(k) for k in ('x','y','w','h')]
            parts.append(f'<polygon id="{ident}" points="{x+w/2:g},{y:g} {x+w:g},{y+h:g} {x:g},{y+h:g}" {style}{transform}/>')
        else:
            raise ValueError('SVG export needs explicit primitive geometry: '+kind)
    for obj in scene['objects']:
        draw(obj)
    result='\n'.join(parts+['</svg>'])+'\n'
    if path is not None:
        path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);path.write_text(result,encoding='utf-8')
    return result


def _transform(value):
    result = Transform()
    for name, text in re.findall(r'([A-Za-z]+)\s*\(([^)]*)\)', value or ''):
        args = [float(x) for x in re.findall(r'[-+]?(?:\d*\.\d+|\d+\.?\d*)(?:[eE][-+]?\d+)?', text)]
        if name == 'translate': matrix = Transform().translate(args[0], args[1] if len(args)>1 else 0)
        elif name == 'scale': matrix = Transform().scale(args[0], args[1] if len(args)>1 else args[0])
        elif name == 'rotate':
            matrix = Transform()
            if len(args)>2: matrix=matrix.translate(args[1],args[2])
            matrix=matrix.rotate(math.radians(args[0]))
            if len(args)>2: matrix=matrix.translate(-args[1],-args[2])
        elif name == 'matrix': matrix = Transform(*args)
        else: raise ValueError('Unsupported transform '+name)
        result = result.transform(matrix)
    return result


def _clip_segment(a,b,box):
    xmin,ymin,xmax,ymax=box
    dx,dy=b[0]-a[0],b[1]-a[1]
    lo,hi=0.,1.
    for p,q in [(-dx,a[0]-xmin),(dx,xmax-a[0]),(-dy,a[1]-ymin),(dy,ymax-a[1])]:
        if abs(p)<1e-14:
            if q<0:return None
        else:
            t=q/p
            if p<0:lo=max(lo,t)
            else:hi=min(hi,t)
            if lo>hi:return None
    return ([a[0]+lo*dx,a[1]+lo*dy],[a[0]+hi*dx,a[1]+hi*dy])


def _clip_polygon(points,box):
    result=points
    for axis,bound,lower in [(0,box[0],True),(0,box[2],False),(1,box[1],True),(1,box[3],False)]:
        original=result;result=[]
        if not original:break
        previous=original[-1]
        pin=previous[axis]>=bound if lower else previous[axis]<=bound
        for current in original:
            cin=current[axis]>=bound if lower else current[axis]<=bound
            if cin!=pin:
                t=(bound-previous[axis])/(current[axis]-previous[axis])
                result.append([previous[j]+t*(current[j]-previous[j]) for j in (0,1)])
            if cin:result.append(current)
            previous,pin=current,cin
    return result


def svg_to_scene(source, width=13.333, height=None):
    """Import existing plot geometry to native primitives and live labels.

    Rectangular clip paths in accepted Matplotlib SVGs are respected for lines;
    no model values are recomputed. Original source SVG remains the chart master.
    """
    parser=ET.XMLParser(target=ET.TreeBuilder(insert_comments=True))
    root=ET.fromstring(Path(source).read_text(encoding='utf-8'),parser=parser)
    values=[float(x) for x in root.get('viewBox','').split()]
    if not values:
        values=[0,0,float(re.sub('[^0-9.]','',root.get('width','980'))),float(re.sub('[^0-9.]','',root.get('height','720')))]
    x0,y0,sw,sh=values
    height=float(height) if height else width*sh/sw
    scale=min(width/sw,height/sh)
    objects=[]
    defs={node.get('id'):node for node in root.iter() if isinstance(node.tag,str) and node.get('id')}
    counter=0
    colors={'black':'#000000','white':'#FFFFFF','none':None,'red':'#FF0000','blue':'#0000FF','gray':'#808080','grey':'#808080'}
    def color(value):
        if value is None:return None
        if value in colors:return colors[value]
        if re.fullmatch('#[0-9a-fA-F]{3}',value):return '#'+''.join(x*2 for x in value[1:])
        return value
    def add(obj):
        nonlocal counter
        counter+=1
        obj['id']=f'obj-{counter:04d}'
        obj['name']=f'{obj["type"]} {counter}: '+str(obj.get('text','plot geometry'))[:70]
        objects.append(obj)
    def style(node,parent):
        values=dict(parent)
        values.update({key:value for key,value in node.attrib.items() if key in {'fill','stroke','stroke-width','font-family','font-size','text-anchor','font-weight','font-style','opacity','fill-opacity','stroke-opacity','clip-path'}})
        values.update(dict(pair.split(':',1) for pair in node.get('style','').split(';') if ':' in pair))
        return {key.strip():value.strip() for key,value in values.items()}
    def number(value,default=0):
        m=re.search(r'[-+]?(?:\d*\.\d+|\d+\.?\d*)(?:[eE][-+]?\d+)?',str(value))
        return float(m.group()) if m else default
    def point(transform,x,y):
        a,b=transform.transformPoint((x,y))
        return [(a-x0)*scale,(b-y0)*scale]
    def textbox(text,x,y,transform,styles,raw_size=None):
        if not text.strip():return
        sx=math.hypot(transform.xx,transform.xy)
        sy=math.hypot(transform.yx,transform.yy)
        raw_size=raw_size or number(styles.get('font-size',12),12)
        size=raw_size*max(sx,sy)*scale*72
        p=point(transform,x,y)
        rotation=math.degrees(math.atan2(transform.xy,transform.xx))
        # Native text baseline converted to a single unwrapped editable box.
        w=max(.12,len(text)*size/72*.59+.08)
        h=max(.1,size/72*1.35)
        anchor=styles.get('text-anchor','start')
        px=p[0]-(w/2 if anchor=='middle' else w-.04 if anchor=='end' else .03)
        py=p[1]-size/72*1.03
        add({'type':'text','x':px,'y':py,'w':w,'h':h,'text':text,'font':'Arial','font_size':max(size,3),'color':color(styles.get('fill','#000000')) or '#000000','valign':'top','margin':0,'word_wrap':False,'rotation':rotation,'bold':styles.get('font-weight')=='bold'})
    def walk(node,transform=Transform(),parent=None,inside_defs=False):
        if not isinstance(node.tag,str):return
        kind=node.tag.rsplit('}',1)[-1]
        if kind in {'defs','metadata','title','desc','clipPath','style'}:return
        if kind=='image':raise ValueError('Raster plot content requires a data-backed native-cell redraw before SVG import')
        styles=style(node,parent or {})
        transform=transform.transform(_transform(node.get('transform')))
        if kind=='use':
            ref=node.get('{http://www.w3.org/1999/xlink}href',node.get('href','')).lstrip('#')
            if ref not in defs:raise ValueError('Missing SVG use '+ref)
            walk(defs[ref],transform.translate(number(node.get('x')),number(node.get('y'))),styles)
            return
        # Matplotlib path-font text retains its exact label in an XML comment.
        if kind=='g' and (node.get('id') or '').startswith('text_'):
            comments=[child.text for child in node if child.tag is ET.Comment]
            actual_text=next((child for child in node.iter() if isinstance(child.tag,str) and child.tag.rsplit('}',1)[-1]=='text'),None)
            if comments and actual_text is None:
                child=next((child for child in node if isinstance(child.tag,str) and child.tag.rsplit('}',1)[-1]=='g'),None)
                if child is not None:
                    tr=transform.transform(_transform(child.get('transform')))
                    label=comments[0].strip()
                    label=re.sub(r'\$|\\(?:mathdefault|mathrm|mathregular)\{([^}]+)\}',lambda m:m.group(1) or '',label)
                    label=label.replace('\\mu','μ').replace('\\times','×').replace('\\Delta','Δ').replace('\\rho','ρ').replace('{','').replace('}','')
                    textbox(label,0,0,tr,styles,raw_size=100 if abs(tr.xx)<.5 else 10)
                    return
        if kind=='text':
            spans=[child for child in node if isinstance(child.tag,str) and child.tag.rsplit('}',1)[-1]=='tspan']
            if spans:
                lastx=number(node.get('x'));lasty=number(node.get('y'))
                for child in spans:
                    lastx=number(child.get('x'),lastx);lasty=number(child.get('y'),lasty)
                    textbox(''.join(child.itertext()),lastx,lasty,transform,style(child,styles))
            else:textbox(''.join(node.itertext()),number(node.get('x')),number(node.get('y')),transform,styles)
            return
        fill=color(styles.get('fill','#000000'))
        stroke=color(styles.get('stroke','none'))
        lw=number(styles.get('stroke-width',1))*scale*72*math.hypot(transform.xx,transform.xy)
        opacity=number(styles.get('opacity',1),1)*number(styles.get('stroke-opacity' if fill is None else 'fill-opacity',1),1)
        shape_style={'fill':fill,'line':stroke,'line_width':lw,'opacity':opacity}
        if styles.get('stroke-dasharray','none')!='none':shape_style['_dash']=styles['stroke-dasharray']
        if kind in {'rect','circle','ellipse'}:
            if kind=='rect':x,y,w,h=[number(node.get(k)) for k in ('x','y','width','height')]
            else:
                rx=number(node.get('r',node.get('rx')));ry=number(node.get('r',node.get('ry')))
                x=number(node.get('cx'))-rx;y=number(node.get('cy'))-ry;w=2*rx;h=2*ry
            p,q=point(transform,x,y),point(transform,x+w,y+h)
            add({'type':'rect' if kind=='rect' else 'ellipse','x':min(p[0],q[0]),'y':min(p[1],q[1]),'w':max(abs(q[0]-p[0]),.0001),'h':max(abs(q[1]-p[1]),.0001),**shape_style})
        elif kind in {'path','line','polyline','polygon'}:
            if kind=='path':data=node.get('d','')
            elif kind=='line':data=f'M {node.get("x1",0)} {node.get("y1",0)} L {node.get("x2",0)} {node.get("y2",0)}'
            else:data='M '+re.sub(r'[,\s]+',' ',node.get('points','')).strip()+(' Z' if kind=='polygon' else '')
            if not data:return
            pen=RecordingPen();parse_path(data,TransformPen(pen,transform))
            paths=[];coords=[];closed=False;current=(0,0)
            for command,args in pen.value:
                if command=='moveTo':
                    if coords:paths.append((coords,closed))
                    coords=[args[0]];current=args[0];closed=False
                elif command=='lineTo':coords.append(args[0]);current=args[0]
                elif command=='curveTo':
                    p0=current;p1,p2,p3=args
                    for i in range(1,13):
                        t=i/12;u=1-t;coords.append(tuple(u**3*p0[j]+3*u*u*t*p1[j]+3*u*t*t*p2[j]+t**3*p3[j] for j in (0,1)))
                    current=p3
                elif command=='qCurveTo':
                    p0=current;p1,p2=args
                    for i in range(1,13):
                        t=i/12;u=1-t;coords.append(tuple(u*u*p0[j]+2*u*t*p1[j]+t*t*p2[j] for j in (0,1)))
                    current=p2
                elif command in {'closePath','endPath'}:
                    closed=command=='closePath';paths.append((coords,closed));coords=[]
            if coords:paths.append((coords,closed))
            for coords,closed in paths:
                if len(coords)<2:continue
                pts=[[(x-x0)*scale,(y-y0)*scale] for x,y in coords]
                clipped=[pts]
                clip_match=re.search(r'url\(#([^)]+)\)',styles.get('clip-path',''))
                if clip_match:
                    clip=defs.get(clip_match.group(1))
                    rect=next((child for child in clip if child.tag.rsplit('}',1)[-1]=='rect'),None) if clip is not None else None
                    if rect is None:raise ValueError('Quantitative SVG requires a supported rectangular clip')
                    ca=point(transform,number(rect.get('x')),number(rect.get('y')))
                    cb=point(transform,number(rect.get('x'))+number(rect.get('width')),number(rect.get('y'))+number(rect.get('height')))
                    box=(min(ca[0],cb[0]),min(ca[1],cb[1]),max(ca[0],cb[0]),max(ca[1],cb[1]))
                    if fill is not None:
                        clipped=[_clip_polygon(pts,box)]
                    else:
                        clipped=[];run=[]
                        for a,b in zip(pts,pts[1:]):
                            segment=_clip_segment(a,b,box)
                            if segment is None:
                                if run:clipped.append(run);run=[]
                            else:
                                start,end=segment
                                if run and math.dist(run[-1],start)>1e-8:clipped.append(run);run=[]
                                if not run:run=[start]
                                run.append(end)
                        if run:clipped.append(run)
                for pts in clipped:
                    if len(pts)<2:continue
                    if len(pts)==2:add({'type':'line','x1':pts[0][0],'y1':pts[0][1],'x2':pts[1][0],'y2':pts[1][1],**{k:v for k,v in shape_style.items() if k!='fill'}})
                    else:
                        add({'type':'polygon','points':pts,'_open':not closed,**shape_style})
        else:
            for child in node:walk(child,transform,styles)
    walk(root)
    scene={'slide':{'width':width,'height':height},'background':'#FFFFFF','objects':objects,'_title':Path(source).stem,'_source_svg':str(Path(source)),'_conversion':'Source vector geometry; curves sampled at 12 points per Bézier segment; native editable labels.'}
    validate_scene(scene)
    return scene
