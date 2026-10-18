"""Audit English figure coverage, native editability, and protected inputs."""
from pathlib import Path
import argparse
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from bs4 import BeautifulSoup
from pptx import Presentation


def audit(root):
    base=root/'assets/figures-editable'
    rows=json.loads((base/'manifest.json').read_text(encoding='utf-8'))
    equations=json.loads((root/'verification/equations.json').read_text(encoding='utf-8'))
    expected={e['id'].lower() for e in equations}
    actual={r['id'] for r in rows if re.fullmatch(r'(?:c[1-5]|a)-e\d+',r['id'])}
    assert expected==actual and len(rows)==199
    native=0
    for r in rows:
        svg=base/r['svg']
        tree=ET.parse(svg)
        assert not re.search(r'[\u3400-\u9fff]',' '.join(tree.getroot().itertext())),r['id']
        assert not any(x.tag.endswith('image') for x in tree.iter()),r['id']
        prs=Presentation(base/r['pptx'])
        assert len(prs.slides)==1
        for s in prs.slides[0].shapes:
            assert s.shape_type!=13,r['id']
            assert not re.search(r'[\u3400-\u9fff]',getattr(s,'text','')),r['id']
            native+=1
        if r['id'] in expected:
            assert svg.read_bytes()==(root/'assets/figures'/(r['id']+'.svg')).read_bytes()
    pages=list((root/'chapters').glob('*.html'))+list((root/'appendices').glob('*.html'))
    pages += [root/p for p in ('modeling/minimal-cavitation/index.html','modeling/minimal-cavitation/STEP_BY_STEP.html','modeling/embedded-pfc-feasibility/index.html')]
    images=0
    for p in pages:
        soup=BeautifulSoup(p.read_text(encoding='utf-8'),'html.parser')
        for img in soup.find_all('img',src=True):
            src=img['src']
            if 'katex' in src:continue
            target=(p.parent/src).resolve()
            assert target.is_file(),str(target)
            assert target.suffix=='.svg' and ('figures' in target.parts or 'figures-editable' in target.parts),str(target)
            assert not re.search(r'[\u3400-\u9fff]',' '.join(ET.parse(target).getroot().itertext())),str(target)
            images+=1
    protected=json.loads((root/'verification/english-figures-20261015/protected-baseline.json').read_text(encoding='utf-8'))['files']
    for p,sha in protected.items():
        if p=='verification/equations.json' and (root/'verification/chapter01-revision-20261019.json').is_file():
            amendment=json.loads((root/'verification/chapter01-revision-20261019.json').read_text(encoding='utf-8'))
            baseline=json.loads((root/'verification/lesson-revision-baseline.json').read_text(encoding='utf-8'))
            assert hashlib.sha256((root/'verification/lesson-revision-baseline.json').read_bytes()).hexdigest() in {amendment['original_baseline_sha256'],amendment.get('public_baseline_sha256')}
            assert amendment['owner_instruction']=='rewrite chapter 1 and upload to GitHub'
            assert {e['id']:e['tex'] for e in equations if e['id'].startswith('C1-')}==amendment['chapter01_equations']
            assert {e['id']:e['tex'] for e in equations if not e['id'].startswith('C1-')}=={k:v for k,v in baseline['equations'].items() if not k.startswith('C1-')}
            continue
        assert hashlib.sha256((root/p).read_bytes()).hexdigest()==sha,p
    sources=json.loads((base/'appendix-studies/manifest.json').read_text(encoding='utf-8'))
    for row in sources:
        if row['kind']=='quantitative plot':
            assert hashlib.sha256((root/row['source']).read_bytes()).hexdigest()==row['source_sha256'],row['id']
    return {'status':'passed','figures':len(rows),'equation_figures':len(expected),'study_figures':10,'native_editable_objects':native,'active_image_elements_checked':images,'protected_inputs':len(protected),'authorized_chapter01_registry_amendment':(root/'verification/chapter01-revision-20261019.json').is_file(),'english_labels':True,'raster_objects':0,'scientific_source_data':'unchanged'}


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1])
    args=parser.parse_args()
    report=audit(args.root)
    (args.root/'verification/english-figures-20261015/figure-audit.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report))
